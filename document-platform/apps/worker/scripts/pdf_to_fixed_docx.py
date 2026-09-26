#!/usr/bin/env python3
"""Create a fixed-position, editable DOCX from a PDF.

The page artwork is rendered after removing only PDF text objects. Editable
VML text boxes are then placed at the original line coordinates. This keeps
images and vector artwork visually stable while making detected text editable
in Microsoft Word and compatible office suites.
"""

from __future__ import annotations

import argparse
import copy
import html
import os
import re
import tempfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt


EMU_PER_POINT = 12700


def safe_font_name(name: str) -> str:
    cleaned = re.sub(r"^[A-Z]{6}\+", "", name or "Arial")
    cleaned = re.sub(r"[^A-Za-z0-9 _-]", "", cleaned)[:80] or "Arial"
    lowered = cleaned.lower()
    # PDF base fonts use PostScript family names that Word may substitute with
    # a serif face. Map them to metrically close Office families explicitly.
    if "helvetica" in lowered:
        return "Arial"
    if "times" in lowered:
        return "Times New Roman"
    if "courier" in lowered:
        return "Courier New"
    return cleaned


def rgb_hex(color: int) -> str:
    return f"{color & 0xFFFFFF:06X}"


def extract_lines(page: fitz.Page) -> list[dict]:
    lines: list[dict] = []
    page_dict = page.get_text("dict", flags=fitz.TEXTFLAGS_TEXT)
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            spans = []
            for span in line.get("spans", []):
                text = span.get("text", "")
                if not text.strip():
                    continue
                spans.append(
                    {
                        "text": text,
                        "font": safe_font_name(span.get("font", "Arial")),
                        "size": max(4.0, min(96.0, float(span.get("size", 11.0)))),
                        "color": rgb_hex(int(span.get("color", 0))),
                        "bold": "bold" in span.get("font", "").lower(),
                        "italic": any(
                            token in span.get("font", "").lower()
                            for token in ("italic", "oblique")
                        ),
                        "bbox": tuple(span.get("bbox", line.get("bbox"))),
                    }
                )
            if not spans:
                continue
            bbox = tuple(line.get("bbox"))
            lines.append({"bbox": bbox, "spans": spans})
    return lines


def add_floating_background(paragraph, image_path: str, width_pt: float, height_pt: float) -> None:
    run = paragraph.add_run()
    inline_shape = run.add_picture(image_path, width=Pt(width_pt), height=Pt(height_pt))
    inline = inline_shape._inline
    inline.tag = qn("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "1",
        "locked": "1",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        inline.set(key, value)

    simple_pos = parse_xml(f'<wp:simplePos {nsdecls("wp")} x="0" y="0"/>')
    position_h = parse_xml(
        f'<wp:positionH {nsdecls("wp")} relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
    )
    position_v = parse_xml(
        f'<wp:positionV {nsdecls("wp")} relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
    )
    wrap_none = parse_xml(f'<wp:wrapNone {nsdecls("wp")}/>')
    inline.insert(0, simple_pos)
    inline.insert(1, position_h)
    inline.insert(2, position_v)

    # ECMA-376 requires the wrapping element to follow wp:extent and the
    # optional wp:effectExtent. LibreOffice repairs the old order silently,
    # while Microsoft Word refuses to open a document whose wp:wrapNone comes
    # before wp:extent.
    extent = inline.find(qn("wp:extent"))
    effect_extent = inline.find(qn("wp:effectExtent"))
    wrap_index = inline.index(effect_extent) + 1 if effect_extent is not None else inline.index(extent) + 1
    inline.insert(wrap_index, wrap_none)


def run_xml(span: dict) -> str:
    text = html.escape(span["text"])
    preserve = ' xml:space="preserve"' if span["text"].startswith(" ") or span["text"].endswith(" ") else ""
    bold = "<w:b/>" if span["bold"] else ""
    italic = "<w:i/>" if span["italic"] else ""
    half_points = max(8, round(span["size"] * 2))
    font = html.escape(span["font"], quote=True)
    return (
        "<w:r><w:rPr>"
        f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}"/>'
        f"{bold}{italic}<w:color w:val=\"{span['color']}\"/>"
        f'<w:sz w:val="{half_points}"/><w:szCs w:val="{half_points}"/>'
        f"</w:rPr><w:t{preserve}>{text}</w:t></w:r>"
    )


def add_editable_textbox(paragraph, line: dict, shape_id: int) -> None:
    x0, y0, x1, y1 = line["bbox"]
    # PDF font metrics and the closest installed Word font can differ slightly.
    # Reserve a small right-side editing area so a line does not wrap merely
    # because Word substituted Helvetica with Arial/Liberation Sans.
    largest_font = max(span["size"] for span in line["spans"])
    width = max(8.0, x1 - x0 + largest_font * 2.5)
    # Office applications reserve additional ascent/descent space inside VML
    # text boxes. A taller box prevents small fonts from being clipped while
    # keeping the editable baseline anchored at the original PDF coordinate.
    top = max(0.0, y0 - 2.0)
    height = max(12.0, (y1 - y0) * 1.8 + 4.0)
    runs = "".join(run_xml(span) for span in line["spans"])
    # Legacy VML text boxes remain the most broadly editable positioned-text
    # representation in desktop Word. The referenced shape type must be
    # declared in the document; without it Word for macOS rejects the whole
    # file even though LibreOffice opens it after an implicit repair.
    shape_type = ""
    if shape_id == 1:
        shape_type = """
            <v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202"
              path="m,l,21600r21600,l21600,xe">
              <v:stroke joinstyle="miter"/>
              <v:path gradientshapeok="t" o:connecttype="rect"/>
            </v:shapetype>
        """
    shape = parse_xml(
        f"""
        <w:r {nsdecls('w')} xmlns:v="urn:schemas-microsoft-com:vml"
          xmlns:o="urn:schemas-microsoft-com:office:office">
          <w:pict>
            {shape_type}
            <v:shape id="AppToolkitText{shape_id}" type="#_x0000_t202"
              o:spid="_x0000_s{1024 + shape_id}"
              style="position:absolute;margin-left:{x0:.3f}pt;margin-top:{top:.3f}pt;width:{width:.3f}pt;height:{height:.3f}pt;z-index:2;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
              stroked="f" filled="f">
              <v:textbox inset="0,0,0,0">
                <w:txbxContent>
                  <w:p>
                    <w:pPr><w:spacing w:before="0" w:after="0"/></w:pPr>
                    {runs}
                  </w:p>
                </w:txbxContent>
              </v:textbox>
            </v:shape>
          </w:pict>
        </w:r>
        """
    )
    paragraph._p.append(shape)


def configure_section(section, width: float, height: float) -> None:
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.top_margin = Pt(0)
    section.right_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def convert(input_path: str, output_path: str, dpi: int) -> None:
    source = fitz.open(input_path)
    background = fitz.open(stream=source.tobytes(garbage=4, deflate=True), filetype="pdf")
    pages: list[dict] = []
    for index, page in enumerate(source):
        lines = extract_lines(page)
        pages.append({"width": page.rect.width, "height": page.rect.height, "lines": lines})
        bg_page = background[index]
        for line in lines:
            for span in line["spans"]:
                rect = fitz.Rect(span["bbox"])
                rect.x0 -= 0.35
                rect.x1 += 0.35
                bg_page.add_redact_annot(rect, fill=False, cross_out=False)
        if lines:
            bg_page.apply_redactions(
                images=fitz.PDF_REDACT_IMAGE_NONE,
                graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                text=fitz.PDF_REDACT_TEXT_REMOVE,
            )

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    shape_id = 1

    with tempfile.TemporaryDirectory(prefix="apptoolkit-fixed-docx-") as temp_dir:
        for index, page_info in enumerate(pages):
            section = document.sections[0] if index == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            configure_section(section, page_info["width"], page_info["height"])
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)

            bg_page = background[index]
            pixmap = bg_page.get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
            image_path = os.path.join(temp_dir, f"page-{index + 1}.png")
            pixmap.save(image_path)
            add_floating_background(
                paragraph, image_path, page_info["width"], page_info["height"]
            )
            for line in page_info["lines"]:
                add_editable_textbox(paragraph, line, shape_id)
                shape_id += 1

        # python-docx leaves its initial empty paragraph before our content.
        body = document._element.body
        for child in list(body):
            if child.tag == qn("w:p") and not child.xpath(".//w:drawing | .//w:pict"):
                body.remove(child)
                break
        document.core_properties.author = "AppToolkitLab"
        document.core_properties.title = "Fixed-position editable PDF export"
        document.save(output_path)

    source.close()
    background.close()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_pdf")
    parser.add_argument("output_docx")
    parser.add_argument("--dpi", type=int, default=180)
    args = parser.parse_args()
    output = Path(args.output_docx)
    output.parent.mkdir(parents=True, exist_ok=True)
    convert(args.input_pdf, str(output), max(96, min(300, args.dpi)))


if __name__ == "__main__":
    main()
